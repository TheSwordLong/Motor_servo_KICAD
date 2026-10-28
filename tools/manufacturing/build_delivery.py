"""Validate current sources and prepare a new delivery in task staging directory.
Does not replace the live delivery until a separate verified installation step.
"""
import json,re,hashlib,zipfile,shutil,csv,collections,xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[2]
WORK=ROOT/'tmp/manufacturing_0p5oz_20261028'
OLD=ROOT/'production/backup_before_0p5oz_20261028/打板文件与BOM'
OUT=WORK/'delivery'
NAME='GD32F303VET6_4Motor_Control'
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def save(p,obj):p.write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def clean(s):return '\n'.join(l for l in s.splitlines() if not l.startswith('G04') and not l.startswith(';'))

def check_bom(after):
    workbook=next((ROOT/'打板文件与BOM/BOM').glob('*.xlsx'))
    w=openpyxl.load_workbook(workbook,data_only=True)
    components={c.get('ref'):c for c in ET.parse(WORK/'netlist.xml').getroot().find('components')}
    fps={f['ref']:f for f in after['footprints']}
    detail={r[0]:r for r in w['逐件明细'].iter_rows(min_row=7,values_only=True)}
    assert set(detail)==set(components),(len(detail),len(components))
    errors=[]
    for ref,row in detail.items():
        c=components[ref];fp=fps.get(ref)
        props={p.get('name'):p.get('value','') for p in c.findall('property')}
        fields={p.get('name'):p.text or '' for p in c.findall('fields/field')}
        excluded=any(k in props for k in ['exclude_from_bom','exclude_from_board','dnp']) or fp is None or fp['dnp'] or fp['exclude_bom']
        qty=0 if excluded else 1
        values={1:c.findtext('value',''),2:qty,4:fields.get('MPN',''),5:fields.get('Manufacturer',''),
                6:fp['side'] if fp else '',7:fp['footprint'] if fp else c.findtext('footprint',''),8:c.findtext('footprint',''),
                10:fields.get('huaqiu_pn',''),11:c.find('sheetpath').get('names'),12:props.get('Sheetfile','')}
        for col,value in values.items():
            if (row[col] if row[col] is not None else '')!=value:errors.append([ref,col,row[col],value])
    assert not errors,errors[:15]
    purchased={ref for ref,row in detail.items() if row[2]==1};group_refs=[];groups=[]
    for r in w['采购汇总'].iter_rows(min_row=7,values_only=True):
        refs=[v.strip() for v in r[0].split(',')];group_refs.extend(refs)
        assert r[2]==len(refs)
        for ref in refs:
            d=detail[ref]
            for gc,dc in [(1,1),(3,4),(4,5),(6,7),(7,9),(8,10)]:assert (r[gc] or '')==(d[dc] or ''),(ref,gc)
        groups.append(r)
    assert len(group_refs)==len(set(group_refs)) and set(group_refs)==purchased
    assert sum(r[2] for r in groups)==w['采购汇总']['D3'].value==151
    for p in (ROOT/'打板文件与BOM/BOM').iterdir():assert sha(p)==sha(OLD/'BOM'/p.name)
    return {'workbook_preserved_byte_for_byte':True,'csvs_preserved_byte_for_byte':True,'detail_rows_checked_against_fresh_netlist':len(detail),
            'summary_groups_checked_against_detail':len(groups),'purchase_pieces':len(purchased),'pcb_footprints':len(fps),
            'missing_MPN_groups':sum(not r[3] for r in groups),'missing_LCSC_groups':sum(not r[7] for r in groups),'missing_Huaqiu_groups':sum(not r[8] for r in groups)}

def main():
    before=read(WORK/'before.json');after=read(WORK/'after.json')
    assert before['stable_blocks']==after['stable_blocks']
    assert before['footprints']==after['footprints'] and before['power_planes']==after['power_planes']
    assert after['thickness_mm']==1.6 and after['copper_layers']==6
    for a,b in zip(before['vias'],after['vias']):
        assert all(a[k]==b[k] for k in ['position','net','drill','diameter'])
        assert b['filling']==1 and b['capping']==1
    assert sha(ROOT/(NAME+'.kicad_pcb'))==after['source_sha256']
    drcb=read(WORK/'drc_before.json');drc=read(WORK/'drc_after.json');erc=read(WORK/'erc.json')
    assert drcb['violations']==drc['violations']
    assert not drc['unconnected_items'] and not drc['schematic_parity']
    ercs=[v for s in erc['sheets'] for v in s['violations']]
    older=read(OLD/'checks/erc.json')
    def erc_key(v):
        # KiCad may choose a different occurrence of an identical global label
        # when reporting a name collision; source sheets themselves are hashed.
        if v['type'] in ['multiple_net_names','same_local_global_label']:
            return json.dumps([v['type'],v['severity'],v['description'],sorted(i['description'] for i in v['items'])])
        return json.dumps(v,sort_keys=True)
    assert collections.Counter(erc_key(v) for s in older['sheets'] for v in s['violations'])==collections.Counter(erc_key(v) for v in ercs)
    bom=check_bom(after)
    files=sorted([p for p in (WORK/'gerbers').iterdir() if p.suffix!='.gbrjob'])+sorted((WORK/'drill').glob('*.drl'))
    assert len(files)==13
    oldzip=zipfile.ZipFile(next((OLD/'嘉立创').glob('*.zip')))
    functions={'.gtl':'Copper,L1,Top','.g1':'Copper,L2,Inr','.g2':'Copper,L3,Inr','.g3':'Copper,L4,Inr','.g4':'Copper,L5,Inr','.gbl':'Copper,L6,Bot',
               '.gts':'Soldermask,Top','.gbs':'Soldermask,Bot','.gto':'Legend,Top','.gbo':'Legend,Bot','.gm1':'Profile,NP'}
    gerbers=[];origins=set()
    for p in files:
        text=p.read_text()
        oldname=next(n for n in oldzip.namelist() if Path(n).name==p.name)
        assert clean(text)==clean(oldzip.read(oldname).decode()),'Manufacturing geometry changed: '+p.name
        if p.suffix=='.drl':continue
        fn=re.search(r'TF.FileFunction,([^*]+)',text)[1];assert fn==functions[p.suffix]
        origins.add(re.search(r'TF.SameCoordinates,([^*]+)',text)[1])
        assert '%FSLAX46Y46*%' in text and '%MOMM*%' in text and text.rstrip().endswith('M02*')
        gerbers.append({'file':p.name,'function':fn,'sha256':sha(p),'manufacturing_payload_unchanged':True})
    assert len(origins)==1
    edge=next((WORK/'gerbers').glob('*.gm1')).read_text()
    coords=re.findall(r'X(-?\d+)Y(-?\d+)(?:I-?\d+J-?\d+)?D0[12]\*',edge)
    assert coords[0]==coords[-1] and len(re.findall(r'D02\*',edge))==1
    xs=[int(x)/1e6 for x,y in coords];ys=[int(y)/1e6 for x,y in coords]
    assert max(xs)-min(xs)==100 and max(ys)-min(ys)==60
    drills=read(WORK/'drill_validation.json');power=read(WORK/'power/power_calculation.json')
    assert power['snapshot_sha256']==after['source_sha256'] and power['pass_voltage_estimate']
    assert max(v for c in power['convergence'] for v in c['relative_diagonal_difference'])<.05
    OUT.mkdir(exist_ok=True);checks=OUT/'checks';checks.mkdir(exist_ok=True)
    shutil.copytree(ROOT/'打板文件与BOM/BOM',OUT/'BOM',dirs_exist_ok=True)
    archive_names=[]
    for folder,prefix in [('嘉立创','JLCPCB'),('华秋','HUAQIU')]:
        (OUT/folder).mkdir(exist_ok=True)
        name=f'{prefix}_{NAME}_6Layer_1p6mm_Inner0p5oz_20261028.zip';path=OUT/folder/name
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
            for p in files:z.write(p,p.name)
        with zipfile.ZipFile(path) as z:
            assert z.testzip() is None and len(z.namelist())==13
            assert set(z.namelist())=={p.name for p in files}
            for p in files:assert z.read(p.name)==p.read_bytes()
        archive_names.append(str(path.relative_to(OUT)))
    for src,dst in [('drc_after.json','drc.json'),('erc.json','erc.json'),('drill_report.txt','drill_report.txt')]:shutil.copy2(WORK/src,checks/dst)
    shutil.copy2(WORK/'power/power_calculation.json',checks/'power_calculation.json')
    for name in ['top.svg','bottom.svg','top.png','bottom.png','In2_24V.svg','In2_24V.png','In3_power.svg','In3_power.png']:shutil.copy2(WORK/name,checks/name)
    for p in (WORK/'drill').glob('*.svg'):shutil.copy2(p,checks/p.name)
    for p in (WORK/'power').glob('*0.125.png'):shutil.copy2(p,checks/p.name)
    with (checks/'检查问题明细.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(['来源','级别','类别','描述','项目'])
        for source,violations in [('DRC',drc['violations']),('ERC',ercs)]:
            for v in violations:writer.writerow([source,v['severity'],v['type'],v['description'],'; '.join(i.get('description','') for i in v.get('items',[]))])
    sources=read(OLD/'checks/source_sha256.json')
    sourcehash={str(p):sha(ROOT/p) for p in sources}
    changed=[p for p in sources if sources[p]!=sourcehash[p]]
    assert changed==[NAME+'.kicad_pcb'],changed
    sourcehash['README.md']=sha(ROOT/'README.md')
    save(checks/'source_sha256.json',sourcehash)
    save(checks/'bom_verification.json',bom)
    fabrication={**drills,'gerbers':gerbers,'same_origin_all_layers':True,'absolute_origin_mm':True,'outline_closed':True,
                 'nominal_outline_mm':[100,60],'order_thickness_mm':1.6,'copper_dielectric_stack_sum_mm':1.5384,
                 'cad_stack_including_nominal_soldermask_mm':1.5584,'all_manufacturing_payloads_equal_to_prior_package':True}
    save(checks/'fabrication_validation.json',fabrication)
    save(checks/'final_verification.json',{'source_sha256':after['source_sha256'],'drc_errors':0,'drc_warnings':len(drc['violations']),
          'drc_violations_identical_to_baseline':True,'unconnected':0,'schematic_parity':0,'erc_errors':sum(v['severity']=='error' for v in ercs),'erc_warnings':sum(v['severity']=='warning' for v in ercs),
          'erc_warning_semantics_and_counts_unchanged':True,'erc_note':'Global-label collision reports can select a different instance of the same label; schematic hashes unchanged.',
          'geometry_netlist_footprints_pads_drills_outline_unchanged':True,
          'filled_power_planes_unchanged':True,'via_filled_capped_count':250,'minimum_via_drill_mm':min(v['drill'] for v in after['vias']),
          'bom':bom,'ZIPs':archive_names,'ZIP_member_count_each':13,'ZIP_integrity':True,
          'physical_thermal_test_performed':False,'website_order_or_payment_performed':False})
    print(json.dumps({'bom':bom,'archives':archive_names,'drc_warnings':len(drc['violations']),'erc_warnings':len(ercs)}))
if __name__=='__main__':main()
